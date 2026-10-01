"""Uploaded files -> text sections (per page for PDFs) -> overlapping chunks."""

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
CHUNK_CHARS = 800
OVERLAP_CHARS = 150


class ParseError(Exception):
    pass


@dataclass
class Section:
    text: str
    page: int | None = None


@dataclass
class Chunk:
    text: str
    page: int | None


def extract(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [Section(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [Section("\n\n".join(paragraphs))]
    if suffix in (".txt", ".md"):
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1254", "latin-1"):  # cp1254: Turkish Windows files
            try:
                return [Section(raw.decode(encoding))]
            except UnicodeDecodeError:
                continue
    raise ParseError(f"Desteklenmeyen dosya türü: {suffix or 'uzantısız'} (desteklenen: PDF, DOCX, TXT, MD)")


def _units(text: str) -> list[str]:
    """Paragraphs, and sentences of paragraphs too long to keep whole."""
    units: list[str] = []
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = re.sub(r"(?m)^\s*#{1,6}\s+", "", paragraph)  # markdown heading markers
        paragraph = re.sub(r"\s+", " ", paragraph).strip()
        if not paragraph:
            continue
        if len(paragraph) <= CHUNK_CHARS:
            units.append(paragraph)
            continue
        sentences = re.split(r"(?<=[.!?])\s+", paragraph)
        for sentence in sentences:
            while len(sentence) > CHUNK_CHARS:  # a run-on "sentence" (tables, lists)
                units.append(sentence[:CHUNK_CHARS])
                sentence = sentence[CHUNK_CHARS - OVERLAP_CHARS :]
            if sentence:
                units.append(sentence)
    return units


def _overlap(units: list[str]) -> list[str]:
    """The end of a chunk to repeat at the start of the next: whole paragraphs, or else whole
    sentences of the last one, so a chunk never starts mid-sentence (people read and edit them)."""
    limit = OVERLAP_CHARS * 2
    kept: list[str] = []
    size = 0
    for unit in reversed(units):
        if size + len(unit) > limit:
            break
        kept.insert(0, unit)
        size += len(unit) + 1
    if kept:
        return kept
    sentences = re.split(r"(?<=[.!?])\s+", units[-1])
    for sentence in reversed(sentences[1:]):  # never the whole paragraph again
        if size + len(sentence) > limit:
            break
        kept.insert(0, sentence)
        size += len(sentence) + 1
    return [" ".join(kept)] if kept else []


def chunk(sections: list[Section]) -> list[Chunk]:
    """Pack paragraphs into ~CHUNK_CHARS pieces; each chunk repeats the end of the previous
    one so an answer split across a boundary is still found whole."""
    chunks: list[Chunk] = []
    for section in sections:
        units_in: list[str] = []
        for unit in _units(section.text):
            if units_in and sum(len(u) + 1 for u in units_in) + len(unit) > CHUNK_CHARS:
                chunks.append(Chunk("\n".join(units_in), section.page))
                units_in = _overlap(units_in)
            units_in.append(unit)  # joined with line breaks to keep paragraphs readable
        if units_in:
            chunks.append(Chunk("\n".join(units_in), section.page))
    return chunks
